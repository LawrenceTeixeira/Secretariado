#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Gera documentos no padrão Secretariado Sob Medida (papel timbrado).

Uso:
  python scripts/gerar_documento.py conteudo.json saida.docx
  python scripts/gerar_documento.py conteudo.json saida.docx --html saida.html
  python scripts/gerar_documento.py --papel-timbrado papel_timbrado.docx
  python scripts/gerar_documento.py --exemplo exemplo.json

O conteúdo vem de um JSON (formato em references/formato_json.md). O visual
(logo, cores, fontes, cabeçalho e rodapé) fica todo aqui, para que qualquer
documento saia igual, não importa quem o peça.

Se a biblioteca python-docx não estiver disponível, o script gera um HTML no
mesmo padrão no lugar do .docx e avisa na saída.
"""
import argparse
import base64
import html
import json
import re
import sys
from pathlib import Path

SKILL_DIR = Path(__file__).resolve().parent.parent
LOGO = SKILL_DIR / "assets" / "logo_horizontal.png"

# Identidade visual (references/identidade_visual.md)
VINHO = "532832"
SALVIA = "C8C59A"
SALVIA_CLARA = "EEEDE1"
ZEBRA = "F7F6EF"
LINHA = "D9D6C3"
TEXTO = "2B2B2B"
CINZA = "6B6B6B"
FONTE_TITULO = "Georgia"
FONTE_TEXTO = "Calibri"
MARCA = "Secretariado Sob Medida"
TAGLINE = "Formação, Networking e Transformação Profissional"
STATUS_PADRAO = "Rascunho para revisão — nada foi enviado, agendado ou contratado."
ROTULOS = {
    "briefing": "BRIEFING",
    "ata": "ATA DE REUNIÃO",
    "email": "RASCUNHO DE E-MAIL",
    "plano": "PLANO DE REUNIÃO",
    "providencias": "QUADRO DE PROVIDÊNCIAS",
    "comparacao": "COMPARAÇÃO DE OPÇÕES",
    "agenda": "PROPOSTA DE AGENDA",
    "livre": "DOCUMENTO",
}
BLOCOS = ("texto", "itens", "tabela", "checklist", "destaque", "nota")

EXEMPLO = {
    "tipo": "ata",
    "titulo": "Ata da reunião do comitê do Fórum Atlas",
    "subtitulo": "Reunião de 02/03/2027",
    "metadados": [["Data", "02/03/2027"], ["Participantes", "Marina Souza, Paulo Lima e Renata Alves"],
                  ["Preparado por", "Agente de IA, para revisão do secretariado"]],
    "secoes": [
        {"titulo": "Decisões", "itens": ["Manter o fórum em 18/03/2027 [T2]."]},
        {"titulo": "Tarefas", "tabela": {"colunas": ["Ação", "Responsável", "Prazo", "Origem"],
                                         "linhas": [["Enviar a lista de convidados", "Renata", "05/03/2027", "T4"]]}},
        {"titulo": "Rascunho de e-mail", "destaque": {"titulo": "Assunto: Fórum Atlas | pendências",
                                                      "texto": "Olá, equipe.\nSegue o resumo para revisão."}},
    ],
    "revisao_humana": ["Confirmar a data da próxima reunião."],
    "fontes": ["transcricao_comite.txt (T1 a T8)"],
    "ficticio": True,
}


# ---------------------------------------------------------------- entrada
def ler_texto(caminho):
    """Lê o JSON em UTF-8, com ou sem BOM, ou em UTF-16 (o que o Windows PowerShell grava com >)."""
    b = Path(caminho).read_bytes()
    if b.startswith((b"\xff\xfe", b"\xfe\xff")):
        return b.decode("utf-16")
    try:
        return b.decode("utf-8-sig")
    except UnicodeDecodeError:
        print("AVISO: o JSON não está em UTF-8; li como Windows-1252. Confira os acentos no documento.", file=sys.stderr)
        return b.decode("cp1252", errors="replace")


def carregar(caminho):
    try:
        dados = json.loads(ler_texto(caminho))
    except json.JSONDecodeError as e:
        raise SystemExit(f"ERRO: o JSON tem um erro de formato na linha {e.lineno}, coluna {e.colno}: {e.msg}.")
    if not isinstance(dados, dict):
        raise SystemExit("ERRO: o JSON precisa ser um objeto com 'titulo' e 'secoes'.")
    if not dados.get("titulo"):
        raise SystemExit("ERRO: falta o campo 'titulo'.")
    if not isinstance(dados.get("secoes", []), list):
        raise SystemExit("ERRO: 'secoes' precisa ser uma lista.")
    if not dados.get("revisao_humana"):
        print("AVISO: 'revisao_humana' está vazio. O padrão pede a lista do que depende de revisão humana.",
              file=sys.stderr)
    return dados


def blocos_da_secao(sec):
    """Devolve os blocos na ordem: 'conteudo' (ordem livre) ou as chaves curtas na ordem padrão."""
    if isinstance(sec.get("conteudo"), list):
        for bloco in sec["conteudo"]:
            for chave in BLOCOS:
                if chave in bloco:
                    yield chave, bloco[chave]
        return
    for chave in BLOCOS:
        if sec.get(chave):
            yield chave, sec[chave]


def secoes_finais(dados, compacto=False):
    secoes = list(dados.get("secoes", []))
    if dados.get("revisao_humana"):
        secoes.append({"titulo": "O que depende de revisão humana", "checklist": dados["revisao_humana"]})
    if dados.get("fontes"):
        if compacto:  # numa linha só, sem título, para economizar espaço
            secoes.append({"nota": "Fontes consultadas: " + "; ".join(map(str, dados["fontes"])) + ".", "_sem_titulo": True})
        else:
            secoes.append({"titulo": "Fontes consultadas", "itens": dados["fontes"], "_pequeno": True})
    return secoes


def como_lista(valor):
    return valor if isinstance(valor, list) else [valor]


def partes_negrito(texto):
    """Divide o texto em trechos (conteúdo, negrito?) usando **negrito**."""
    partes = re.split(r"(\*\*.+?\*\*)", str(texto))
    for p in partes:
        if p.startswith("**") and p.endswith("**") and len(p) > 4:
            yield p[2:-2], True
        elif p:
            yield p, False


# ---------------------------------------------------------------- DOCX
def gerar_docx(dados, saida, papel_timbrado=False):
    from docx import Document
    from docx.enum.section import WD_ORIENT
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor

    def cor(hexa):
        return RGBColor.from_string(hexa)

    def fonte(run, nome=FONTE_TEXTO, tamanho=None, negrito=None, italico=None, cor_hex=None, versalete=False):
        run.font.name = nome
        rpr = run._element.get_or_add_rPr()
        rfonts = rpr.find(qn("w:rFonts"))
        if rfonts is None:
            rfonts = OxmlElement("w:rFonts")
            rpr.insert(0, rfonts)
        for att in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
            rfonts.set(qn(att), nome)
        if tamanho:
            run.font.size = Pt(tamanho)
        if negrito is not None:
            run.font.bold = negrito
        if italico is not None:
            run.font.italic = italico
        if cor_hex:
            run.font.color.rgb = cor(cor_hex)
        if versalete:
            run.font.small_caps = True
        return run

    def texto_rico(par, texto, tamanho=None, cor_hex=None, base_negrito=False):
        for trecho, negrito in partes_negrito(texto):
            fonte(par.add_run(trecho), tamanho=tamanho, negrito=negrito or base_negrito, cor_hex=cor_hex)

    def espaco(par, antes=0, depois=4, linha=1.15):
        pf = par.paragraph_format
        pf.space_before = Pt(antes)
        pf.space_after = Pt(depois)
        pf.line_spacing = linha

    def borda_paragrafo(par, lado="bottom", tamanho=8, cor_hex=VINHO, espaco_pt=4):
        ppr = par._p.get_or_add_pPr()
        pbdr = ppr.find(qn("w:pBdr"))
        if pbdr is None:
            pbdr = OxmlElement("w:pBdr")
            ppr.append(pbdr)
        b = OxmlElement(f"w:{lado}")
        b.set(qn("w:val"), "single")
        b.set(qn("w:sz"), str(tamanho))
        b.set(qn("w:space"), str(espaco_pt))
        b.set(qn("w:color"), cor_hex)
        pbdr.append(b)

    def sombrear(celula, hexa):
        tcpr = celula._tc.get_or_add_tcPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:color"), "auto")
        shd.set(qn("w:fill"), hexa)
        tcpr.append(shd)

    def margens_celula(celula, topo=60, lado=110):
        tcpr = celula._tc.get_or_add_tcPr()
        mar = OxmlElement("w:tcMar")
        for nome, valor in (("top", topo), ("bottom", topo), ("start", lado), ("end", lado)):
            m = OxmlElement(f"w:{nome}")
            m.set(qn("w:w"), str(valor))
            m.set(qn("w:type"), "dxa")
            mar.append(m)
        tcpr.append(mar)

    def bordas_tabela(tabela, cor_hex=LINHA, tamanho=4, internas=True, nenhuma=False):
        tblpr = tabela._tbl.tblPr
        bordas = OxmlElement("w:tblBorders")
        for lado in ("top", "left", "bottom", "right", "insideH", "insideV"):
            b = OxmlElement(f"w:{lado}")
            if nenhuma or (not internas and lado.startswith("inside")):
                b.set(qn("w:val"), "nil")
            else:
                b.set(qn("w:val"), "single")
                b.set(qn("w:sz"), str(tamanho))
                b.set(qn("w:color"), cor_hex)
            bordas.append(b)
        tblpr.append(bordas)

    def largura_total(tabela):
        tblpr = tabela._tbl.tblPr
        w = tblpr.find(qn("w:tblW"))
        if w is None:
            w = OxmlElement("w:tblW")
            tblpr.append(w)
        w.set(qn("w:type"), "pct")
        w.set(qn("w:w"), "5000")

    def borda_celula(celula, lado, tamanho, cor_hex):
        tcpr = celula._tc.get_or_add_tcPr()
        bordas = tcpr.find(qn("w:tcBorders"))
        if bordas is None:
            bordas = OxmlElement("w:tcBorders")
            tcpr.append(bordas)
        b = OxmlElement(f"w:{lado}")
        b.set(qn("w:val"), "single")
        b.set(qn("w:sz"), str(tamanho))
        b.set(qn("w:color"), cor_hex)
        bordas.append(b)

    def repetir_cabecalho(linha):
        trpr = linha._tr.get_or_add_trPr()
        th = OxmlElement("w:tblHeader")
        th.set(qn("w:val"), "true")
        trpr.append(th)

    def nao_quebrar_linha(linha):
        trpr = linha._tr.get_or_add_trPr()
        c = OxmlElement("w:cantSplit")
        c.set(qn("w:val"), "true")
        trpr.append(c)

    def campo(par, instrucao, tamanho, cor_hex):
        for tipo, texto in (("begin", None), (None, instrucao), ("separate", None), (None, "1"), ("end", None)):
            run = fonte(par.add_run(), tamanho=tamanho, cor_hex=cor_hex)
            if tipo:
                fc = OxmlElement("w:fldChar")
                fc.set(qn("w:fldCharType"), tipo)
                run._r.append(fc)
            elif texto == instrucao:
                it = OxmlElement("w:instrText")
                it.set(qn("xml:space"), "preserve")
                it.text = f" {instrucao} "
                run._r.append(it)
            else:
                run.text = texto

    doc = Document()

    # Página A4 e margens
    sec = doc.sections[0]
    sec.orientation = WD_ORIENT.PORTRAIT
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(2.0)
    sec.top_margin, sec.bottom_margin = (Cm(2.3), Cm(1.9)) if dados.get("compacto") else (Cm(2.6), Cm(2.2))
    sec.header_distance, sec.footer_distance = Cm(0.9), Cm(0.9)

    # Estilos de base (editáveis no Word)
    normal = doc.styles["Normal"]
    normal.font.name = FONTE_TEXTO
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = cor(TEXTO)
    rpr = normal.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for att in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rfonts.set(qn(att), FONTE_TEXTO)
    lang = OxmlElement("w:lang")
    lang.set(qn("w:val"), "pt-BR")
    lang.set(qn("w:eastAsia"), "pt-BR")
    rpr.append(lang)
    normal.paragraph_format.space_after = Pt(4)
    normal.paragraph_format.line_spacing = 1.15

    rotulo = dados.get("rotulo") or ROTULOS.get(dados.get("tipo", "livre"), "DOCUMENTO")
    # Modo compacto: mesmo visual com menos espaço vertical, para caber em uma página
    compacto = bool(dados.get("compacto"))
    SEC_ANTES, SEC_TAM, TXT_DEPOIS, ITEM_DEPOIS, LINHA_TXT = (5, 12, 2, 0, 1.0) if compacto else (14, 13, 5, 2, 1.15)
    CEL = (25, 90) if compacto else (60, 110)

    # Cabeçalho: logo à esquerda, tipo do documento à direita, filete vinho
    cab = sec.header
    cab.is_linked_to_previous = False
    tab = cab.add_table(rows=1, cols=2, width=Cm(17.0))
    bordas_tabela(tab, nenhuma=True)
    largura_total(tab)
    c_logo, c_rot = tab.rows[0].cells
    p = c_logo.paragraphs[0]
    if LOGO.exists():
        p.add_run().add_picture(str(LOGO), height=Cm(1.15))
    else:
        fonte(p.add_run(MARCA.upper()), FONTE_TITULO, 14, cor_hex=VINHO)
    p = c_rot.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    fonte(p.add_run(rotulo), FONTE_TITULO, 9.5, cor_hex=VINHO, versalete=False)
    if dados.get("codigo"):
        p2 = c_rot.add_paragraph()
        p2.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        fonte(p2.add_run(str(dados["codigo"])), tamanho=8.5, cor_hex=CINZA)
    for c in (c_logo, c_rot):
        c.vertical_alignment = 1  # centro
    filete = cab.paragraphs[0] if cab.paragraphs and not cab.paragraphs[0].text else cab.add_paragraph()
    espaco(filete, 0, 0, 1.0)
    borda_paragrafo(filete, "bottom", 8, VINHO, 1)
    # move o filete para depois da tabela
    tab._tbl.addnext(filete._p)

    # Rodapé: marca e tagline, aviso de IA e paginação
    rod = sec.footer
    rod.is_linked_to_previous = False
    p = rod.paragraphs[0]
    borda_paragrafo(p, "top", 4, SALVIA, 4)
    espaco(p, 0, 1, 1.0)
    fonte(p.add_run(MARCA.upper()), FONTE_TITULO, 8, cor_hex=VINHO)
    fonte(p.add_run("  ·  " + TAGLINE), tamanho=8, cor_hex=CINZA)
    p = rod.add_paragraph()
    espaco(p, 0, 0, 1.0)
    aviso = "" if papel_timbrado else "Preparado com apoio de agente de IA · rascunho para revisão humana"
    if dados.get("ficticio"):
        aviso += " · demonstração com dados fictícios"
    fonte(p.add_run((aviso + "   ·   " if aviso else "") + "página "), tamanho=8, cor_hex=CINZA)
    campo(p, "PAGE", 8, CINZA)
    fonte(p.add_run(" de "), tamanho=8, cor_hex=CINZA)
    campo(p, "NUMPAGES", 8, CINZA)

    # Bloco de título
    p = doc.paragraphs[0] if doc.paragraphs else doc.add_paragraph()
    espaco(p, 2, 2, 1.0)
    fonte(p.add_run(dados["titulo"]), FONTE_TITULO, 18 if compacto else 20, cor_hex=VINHO)
    if dados.get("subtitulo"):
        p = doc.add_paragraph()
        espaco(p, 0, 3 if compacto else 6, 1.0)
        fonte(p.add_run(dados["subtitulo"]), tamanho=11, cor_hex=CINZA)

    # Faixa de status
    status = dados.get("status", STATUS_PADRAO)
    if status:
        t = doc.add_table(rows=1, cols=1)
        bordas_tabela(t, nenhuma=True)
        largura_total(t)
        c = t.rows[0].cells[0]
        sombrear(c, SALVIA_CLARA)
        borda_celula(c, "left", 24, VINHO)
        margens_celula(c, 40 if compacto else 70, 140)
        pc = c.paragraphs[0]
        espaco(pc, 0, 0, 1.0)
        fonte(pc.add_run(status), tamanho=9.5, negrito=True, cor_hex=VINHO)

    # Metadados
    meta = [m for m in dados.get("metadados", []) if isinstance(m, (list, tuple)) and len(m) == 2]
    if meta:
        espacador = doc.add_paragraph()
        espaco(espacador, 0, 0, 0.3 if compacto else 0.6)
        t = doc.add_table(rows=len(meta), cols=2)
        bordas_tabela(t, LINHA, 4)
        largura_total(t)
        for i, (rot, val) in enumerate(meta):
            c1, c2 = t.rows[i].cells
            c1.width, c2.width = Cm(4.2), Cm(12.8)
            sombrear(c1, SALVIA_CLARA)
            for c in (c1, c2):
                margens_celula(c, 20 if compacto else 40, 100)
            pc = c1.paragraphs[0]
            espaco(pc, 0, 0, 1.0)
            fonte(pc.add_run(str(rot)), tamanho=9.5, negrito=True, cor_hex=VINHO)
            pc = c2.paragraphs[0]
            espaco(pc, 0, 0, 1.0)
            texto_rico(pc, val, tamanho=9.5)

    # Seções
    for secao in secoes_finais(dados, compacto):
        pequeno = secao.get("_pequeno", False)
        if not secao.get("_sem_titulo"):
            h = doc.add_paragraph()
            espaco(h, SEC_ANTES, 3 if compacto else 5, 1.0)
            h.paragraph_format.keep_with_next = True
            borda_paragrafo(h, "bottom", 6, SALVIA, 2)
            fonte(h.add_run(str(secao.get("titulo", ""))), FONTE_TITULO, SEC_TAM if not pequeno else 11, cor_hex=VINHO)

        for chave, valor in blocos_da_secao(secao):
            if chave == "texto":
                for paragrafo in como_lista(valor):
                    p = doc.add_paragraph()
                    espaco(p, 0, TXT_DEPOIS, LINHA_TXT)
                    texto_rico(p, paragrafo, tamanho=9 if pequeno else None)
            elif chave in ("itens", "checklist"):
                for item in como_lista(valor):
                    p = doc.add_paragraph(style="List Bullet" if chave == "itens" else None)
                    espaco(p, 0, ITEM_DEPOIS, LINHA_TXT)
                    if chave == "checklist":
                        p.paragraph_format.left_indent = Cm(0.6)
                        p.paragraph_format.first_line_indent = Cm(-0.6)
                        fonte(p.add_run("☐  "), "Segoe UI Symbol", 10, cor_hex=VINHO)
                    texto_rico(p, item, tamanho=9 if pequeno else None)
            elif chave == "tabela":
                colunas = [str(x) for x in valor.get("colunas", [])]
                linhas = [[("" if x is None else str(x)) for x in l] for l in valor.get("linhas", [])]
                if not colunas:
                    continue
                t = doc.add_table(rows=1 + len(linhas), cols=len(colunas))
                t.alignment = WD_TABLE_ALIGNMENT.CENTER
                bordas_tabela(t, LINHA, 4)
                largura_total(t)
                # larguras proporcionais ao conteúdo
                # (colunas curtas recebem largura para não quebrar "não informado" ou datas)
                pesos = []
                for j, col in enumerate(colunas):
                    tam = [len(col)] + [len(l[j].replace("**", "")) if j < len(l) else 0 for l in linhas]
                    pesos.append(max(6, min(60, sum(tam) / len(tam) * 0.6 + max(tam) * 0.4), min(max(tam), 18) * 1.25))
                total = sum(pesos)
                larguras = [Cm(17.0 * w / total) for w in pesos]
                corpo_tam = 9.5 if len(colunas) <= 4 else 9
                for j, col in enumerate(colunas):
                    c = t.rows[0].cells[j]
                    c.width = larguras[j]
                    sombrear(c, VINHO)
                    margens_celula(c, *CEL)
                    pc = c.paragraphs[0]
                    espaco(pc, 0, 0, 1.0)
                    fonte(pc.add_run(col), tamanho=corpo_tam, negrito=True, cor_hex="FFFFFF")
                repetir_cabecalho(t.rows[0])
                for i, linha in enumerate(linhas, start=1):
                    nao_quebrar_linha(t.rows[i])
                    for j in range(len(colunas)):
                        c = t.rows[i].cells[j]
                        c.width = larguras[j]
                        if i % 2 == 0:
                            sombrear(c, ZEBRA)
                        margens_celula(c, *CEL)
                        pc = c.paragraphs[0]
                        espaco(pc, 0, 0, 1.05)
                        texto_rico(pc, linha[j] if j < len(linha) else "", tamanho=corpo_tam)
                if valor.get("legenda"):
                    p = doc.add_paragraph()
                    espaco(p, 1 if compacto else 2, 2 if compacto else 4)
                    fonte(p.add_run(str(valor["legenda"])), tamanho=8.5, italico=True, cor_hex=CINZA)
                else:
                    espaco(doc.add_paragraph(), 0, 0 if compacto else 2, 0.5 if compacto else 0.8)
            elif chave == "destaque":
                t = doc.add_table(rows=1, cols=1)
                bordas_tabela(t, LINHA, 4)
                largura_total(t)
                c = t.rows[0].cells[0]
                sombrear(c, "FBFAF6")
                borda_celula(c, "left", 24, VINHO)
                margens_celula(c, 60 if compacto else 110, 180)
                pc = c.paragraphs[0]
                espaco(pc, 0, 4, 1.0)
                if valor.get("titulo"):
                    texto_rico(pc, valor["titulo"], tamanho=10, cor_hex=VINHO, base_negrito=True)
                    pc = c.add_paragraph()
                    espaco(pc, 0, 4)
                linhas_txt = str(valor.get("texto", "")).split("\n")
                for k, linha in enumerate(linhas_txt):
                    if k:
                        pc = c.add_paragraph()
                        espaco(pc, 0, 4)
                    texto_rico(pc, linha, tamanho=10)
                espaco(doc.add_paragraph(), 0, 0 if compacto else 2, 0.5 if compacto else 0.8)
            elif chave == "nota":
                p = doc.add_paragraph()
                espaco(p, 2, 4)
                fonte(p.add_run(str(valor)), tamanho=8.5, italico=True, cor_hex=CINZA)

    if papel_timbrado:
        p = doc.add_paragraph()
        espaco(p, 10, 4)
        fonte(p.add_run("Escreva aqui o conteúdo. Use os títulos de seção em Georgia vinho e o texto em Calibri."),
              tamanho=10.5, cor_hex=CINZA, italico=True)

    props = doc.core_properties
    props.title = dados["titulo"]
    props.subject = rotulo
    props.author = MARCA
    props.comments = "Gerado pela skill padrao-secretariado-sob-medida. Rascunho para revisão humana."
    props.language = "pt-BR"
    Path(saida).parent.mkdir(parents=True, exist_ok=True)
    doc.save(saida)


# ---------------------------------------------------------------- HTML
def gerar_html(dados, saida):
    E = html.escape

    def rico(t):
        return "".join(f"<strong>{E(x)}</strong>" if b else E(x) for x, b in partes_negrito(t))

    logo = ""
    if LOGO.exists():
        b64 = base64.b64encode(LOGO.read_bytes()).decode("ascii")
        logo = f'<img class="logo" alt="{E(MARCA)}" src="data:image/png;base64,{b64}">'
    else:
        logo = f'<span class="marca">{E(MARCA.upper())}</span>'
    rotulo = dados.get("rotulo") or ROTULOS.get(dados.get("tipo", "livre"), "DOCUMENTO")
    partes = [f'<header><div>{logo}</div><div class="rotulo">{E(rotulo)}</div></header>',
              f'<h1>{E(dados["titulo"])}</h1>']
    if dados.get("subtitulo"):
        partes.append(f'<p class="sub">{E(dados["subtitulo"])}</p>')
    status = dados.get("status", STATUS_PADRAO)
    if status:
        partes.append(f'<div class="status">{E(status)}</div>')
    meta = [m for m in dados.get("metadados", []) if isinstance(m, (list, tuple)) and len(m) == 2]
    if meta:
        partes.append('<table class="meta">' + "".join(
            f"<tr><th>{E(str(a))}</th><td>{rico(b)}</td></tr>" for a, b in meta) + "</table>")
    for secao in secoes_finais(dados):
        partes.append(f'<h2>{E(str(secao.get("titulo", "")))}</h2>')
        for chave, valor in blocos_da_secao(secao):
            if chave == "texto":
                partes += [f"<p>{rico(x)}</p>" for x in como_lista(valor)]
            elif chave == "itens":
                partes.append("<ul>" + "".join(f"<li>{rico(x)}</li>" for x in como_lista(valor)) + "</ul>")
            elif chave == "checklist":
                partes.append('<ul class="check">' + "".join(f"<li>{rico(x)}</li>" for x in como_lista(valor)) + "</ul>")
            elif chave == "tabela":
                cab = "".join(f"<th>{E(str(c))}</th>" for c in valor.get("colunas", []))
                corpo = "".join("<tr>" + "".join(f"<td>{rico('' if c is None else c)}</td>" for c in l) + "</tr>"
                                for l in valor.get("linhas", []))
                partes.append(f'<div class="tw"><table class="dados"><thead><tr>{cab}</tr></thead><tbody>{corpo}</tbody></table></div>')
                if valor.get("legenda"):
                    partes.append(f'<p class="nota">{E(str(valor["legenda"]))}</p>')
            elif chave == "destaque":
                tit = f'<p class="dt">{rico(valor["titulo"])}</p>' if valor.get("titulo") else ""
                corpo = "".join(f"<p>{rico(x)}</p>" for x in str(valor.get("texto", "")).split("\n"))
                partes.append(f'<div class="destaque">{tit}{corpo}</div>')
            elif chave == "nota":
                partes.append(f'<p class="nota">{E(str(valor))}</p>')
    aviso = "Preparado com apoio de agente de IA · rascunho para revisão humana"
    if dados.get("ficticio"):
        aviso += " · demonstração com dados fictícios"
    partes.append(f'<footer><span class="marca">{E(MARCA.upper())}</span> · {E(TAGLINE)}<br>{E(aviso)}</footer>')
    css = f"""
:root {{ --vinho:#{VINHO}; --salvia:#{SALVIA}; --salvia-clara:#{SALVIA_CLARA}; --zebra:#{ZEBRA}; --linha:#{LINHA}; --texto:#{TEXTO}; --cinza:#{CINZA}; }}
* {{ box-sizing:border-box; }}
body {{ margin:0; background:#f2f1ea; color:var(--texto); font:16px/1.5 {FONTE_TEXTO}, "Segoe UI", Arial, sans-serif; }}
main {{ max-width:860px; margin:24px auto; background:#fff; padding:28px 32px 20px; box-shadow:0 1px 4px rgba(0,0,0,.08); }}
header {{ display:flex; justify-content:space-between; align-items:center; gap:12px; border-bottom:2px solid var(--vinho); padding-bottom:10px; }}
.logo {{ height:46px; width:auto; }}
.rotulo {{ font-family:{FONTE_TITULO}, serif; color:var(--vinho); font-size:13px; letter-spacing:.06em; text-align:right; }}
h1 {{ font-family:{FONTE_TITULO}, serif; color:var(--vinho); font-weight:normal; font-size:30px; margin:20px 0 2px; }}
.sub {{ color:var(--cinza); margin:0 0 10px; }}
.status {{ background:var(--salvia-clara); border-left:6px solid var(--vinho); color:var(--vinho); font-weight:bold; font-size:14px; padding:8px 12px; margin:6px 0 12px; }}
h2 {{ font-family:{FONTE_TITULO}, serif; color:var(--vinho); font-weight:normal; font-size:20px; border-bottom:1.5px solid var(--salvia); padding-bottom:3px; margin:22px 0 8px; }}
table {{ border-collapse:collapse; width:100%; }}
.meta th, .meta td, .dados th, .dados td {{ border:1px solid var(--linha); padding:6px 9px; text-align:left; vertical-align:top; font-size:14px; }}
.meta th {{ background:var(--salvia-clara); color:var(--vinho); width:28%; }}
.dados th {{ background:var(--vinho); color:#fff; }}
.dados tbody tr:nth-child(even) td {{ background:var(--zebra); }}
.tw {{ overflow-x:auto; margin:4px 0 8px; }}
.destaque {{ background:#FBFAF6; border:1px solid var(--linha); border-left:6px solid var(--vinho); padding:10px 14px; margin:6px 0 10px; }}
.destaque p {{ margin:0 0 6px; }}
.dt {{ color:var(--vinho); font-weight:bold; }}
ul.check {{ list-style:none; padding-left:4px; }}
ul.check li::before {{ content:"☐  "; color:var(--vinho); }}
.nota {{ color:var(--cinza); font-style:italic; font-size:13px; }}
footer {{ margin-top:26px; border-top:1px solid var(--salvia); padding-top:6px; font-size:12px; color:var(--cinza); }}
footer .marca {{ font-family:{FONTE_TITULO}, serif; color:var(--vinho); }}
@media (max-width:640px) {{ main {{ margin:0; padding:18px 16px; }} h1 {{ font-size:24px; }} }}
@media print {{ body {{ background:#fff; }} main {{ box-shadow:none; margin:0; }} }}
"""
    pagina = (f'<!doctype html><html lang="pt-BR"><head><meta charset="utf-8">'
              f'<meta name="viewport" content="width=device-width, initial-scale=1">'
              f'<title>{E(dados["titulo"])}</title><style>{css}</style></head><body><main>'
              + "\n".join(partes) + "</main></body></html>")
    Path(saida).parent.mkdir(parents=True, exist_ok=True)
    Path(saida).write_text(pagina, encoding="utf-8")


# ---------------------------------------------------------------- CLI
def saida_utf8():
    """Acentos certos no terminal do Windows e quando a saída vai para um arquivo (>)."""
    for fluxo in (sys.stdout, sys.stderr):
        try:
            fluxo.reconfigure(encoding="utf-8")
        except (AttributeError, ValueError):
            pass


def main():
    saida_utf8()
    ap = argparse.ArgumentParser(description="Gera documentos no padrão Secretariado Sob Medida.")
    ap.add_argument("entrada", nargs="?", help="arquivo JSON com o conteúdo")
    ap.add_argument("saida", nargs="?", help="arquivo .docx de saída")
    ap.add_argument("--html", help="também gera uma versão HTML neste caminho")
    ap.add_argument("--papel-timbrado", metavar="SAIDA", help="gera o papel timbrado em branco (.docx)")
    ap.add_argument("--exemplo", nargs="?", const="-", metavar="ARQUIVO",
                    help="grava um JSON de exemplo no arquivo indicado (sem arquivo, mostra na tela)")
    a = ap.parse_args()

    if a.exemplo:
        texto = json.dumps(EXEMPLO, ensure_ascii=False, indent=2)
        if a.exemplo == "-":
            print(texto)
        else:
            Path(a.exemplo).write_text(texto + "\n", encoding="utf-8")
            print(f"OK: exemplo salvo em {a.exemplo}")
        return
    if a.papel_timbrado:
        try:
            import docx  # noqa: F401
        except ImportError:
            raise SystemExit("ERRO: para gerar o papel timbrado, instale o python-docx (python -m pip install python-docx).")
        gerar_docx({"titulo": "Título do documento", "subtitulo": "Subtítulo ou data", "tipo": "livre",
                    "status": "", "secoes": []}, a.papel_timbrado, papel_timbrado=True)
        print(f"OK: papel timbrado salvo em {a.papel_timbrado}")
        return
    if not (a.entrada and a.saida):
        ap.error("informe o JSON de entrada e o .docx de saída (ou use --exemplo / --papel-timbrado)")

    dados = carregar(a.entrada)
    try:
        import docx  # noqa: F401
    except ImportError:
        alternativo = str(Path(a.saida).with_suffix(".html"))
        gerar_html(dados, alternativo)
        print(f"AVISO: python-docx não está disponível neste ambiente. Gerei o HTML no padrão em {alternativo}. "
              "Diga ao usuário que o Word não pôde ser gerado aqui.")
        return
    gerar_docx(dados, a.saida)
    print(f"OK: documento salvo em {a.saida}")
    if a.html:
        gerar_html(dados, a.html)
        print(f"OK: versão HTML salva em {a.html}")


if __name__ == "__main__":
    main()
